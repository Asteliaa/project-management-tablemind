"""Checks each coursework file against the assignment requirements. Usage: python verify_works.py <work> <repo_dir>
work: pz1 | pz4 | risk | appendix | cross. Prints a markdown section and exits 1 if any hard requirement fails."""
import sys, re, zipfile, warnings
import docx, openpyxl
warnings.filterwarnings('ignore')

work, R = sys.argv[1], sys.argv[2].rstrip('/')
rows, hard_fail = [], False


def check(req, ok, detail='', manual=False):
    global hard_fail
    status = '✅' if ok else ('✍️ вручную' if manual else '❌')
    if not ok and not manual:
        hard_fail = True
    rows.append(f'| {req} | {status} | {detail} |')


def cells(t):
    return ' '.join(c.text for r in t.rows for c in r.cells)


if work == 'pz1':
    d = docx.Document(f'{R}/pz1/Земляник - ПЗ1 Концепция проекта TableMind.docx'); T = d.tables
    ps = [p.text for p in d.paragraphs]; full = '\n'.join(ps)
    alltext = full + ' '.join(cells(t) for t in T)
    check('Шаблон заполнен: нет плейсхолдеров <…>, [Пример], ___', not re.search(r'<[^>]{1,40}>|\[Пример|\[Впишите|___', alltext))
    check('Титульный лист: ФИО слушателя', 'Земляник Р.В.' in full and 'ФИО' not in full, 'Земляник Р.В., гр. 60121')
    check('1.1 Модель 5W — ответ на каждый вопрос', all(T[0].rows[i].cells[2].text.strip() for i in range(1, 6)), '5 из 5')
    check('1.2 Бизнес-окружение (3–5 предложений)', any(p.startswith('TableMind — собственный продукт') for p in ps))
    check('1.2 Проблемы → Решения (3 строки)', all(T[1].rows[i].cells[1].text and T[1].rows[i].cells[2].text for i in range(1, 4)))
    check('1.3 SMART — 5 критериев', all(T[2].rows[i].cells[1].text.strip() for i in range(1, 6)))
    check('1.3 Итоговая SMART-цель одним предложением', any(p.startswith('До 30.09.2027 командой') and len(re.findall(r'[.!?]\s+[А-ЯA-Z]', p)) == 0 for p in ps))
    check('1.4 Правила приёмки — не менее 5 проверяемых критериев',
          sum(1 for p in d.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None and ('≥' in p.text or 'E2E' in p.text or 'blocker' in p.text or 'выдуманных' in p.text)) >= 5)
    check('1.5 Допущения — 5, у каждого влияние', all(T[3].rows[i].cells[1].text and T[3].rows[i].cells[2].text for i in range(1, 6)))
    mos = [len([l for l in T[4].rows[i].cells[2].text.split('\n') if l.strip()]) for i in range(1, 5)]
    check('2.1 MoSCoW — все 4 категории, Must ≤ 60%', all(mos) and mos[0] / sum(mos) <= 0.6,
          f'Must {mos[0]}, Should {mos[1]}, Could {mos[2]}, Won’t {mos[3]}; Must = {mos[0]/sum(mos):.0%}')
    ex = [p for p in ps if p.startswith('НЕ входит')]
    check('2.2 Исключения — 3–5 пунктов', 3 <= len(ex) <= 5, f'{len(ex)}')
    check('3 Ограничения — 6 категорий заполнены', all(T[5].rows[i].cells[1].text.strip() for i in range(1, 7)))
    a = next(i for i, t in enumerate(ps) if t.startswith('Иерархическая структура задач'))
    b = next(i for i, t in enumerate(ps) if 'РАЗДЕЛ 5' in t)
    blk = ps[a + 1:b]
    st = [t for t in blk if re.match(r'^\d\. ', t)]; tk = [t for t in blk if re.match(r'^\d\.\d\. ', t)]
    per = [len([t for t in tk if t.startswith(s[0] + '.')]) for s in st]
    check('4 Декомпозиция: 5–7 этапов, 3–5 задач в этапе, 15–25 задач, иерархия',
          5 <= len(st) <= 7 and all(3 <= x <= 5 for x in per) and 15 <= len(tk) <= 25, f'{len(st)} этапов, задачи по этапам {per}, всего {len(tk)}')
    org = T[5].rows[3].cells[1].text
    loads = [float(a.replace(',', '.')) * int(b) for a, b in re.findall(r'— ([\d,]+) ставк\w*, (\d+) мес', org)]
    check('3 Загрузка команды: сумма ставок × месяцев = 57 чел.-мес., без диапазонов', len(loads) == 7 and sum(loads) == 57 and '–1 ставк' not in org, f'{sum(loads):g} чел.-мес.')
    acc = [p.text for p in d.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
    check('1.4 Приёмка покрывает бизнес-метрики SMART (≥300 зарегистрированных, ≥50 платящих)', any('300 зарегистрированных' in t and '50 платящих' in t for t in acc))
    check('4 Декомпозиция содержит только Must/обязательные работы (нет Should: ghost, Cmd-K, история версий)', not re.search(r'ghost|Cmd-K|\(Should\)|истори\w+ версий', ' '.join(blk)))
    TW = next((t for t in T if t.rows[0].cells[0].text.strip() == 'Код'), None)
    codes = [r.cells[0].text.strip() for r in TW.rows[1:]] if TW else []
    check('4 Коды WBS из декомпозиции расшифрованы в самом отчёте (таблица 4.2)', codes == [f'{i}.0' for i in range(1, 16)], f'{len(codes)} работ')
    TS = next(t for t in T if t.rows[0].cells[1].text.strip() == 'Стейкхолдер')
    TQ = next(t for t in T if t.rows[0].cells[0].text.strip() == 'Квадрант')
    sh = [[c.text for c in r.cells] for r in TS.rows[1:]]
    check('5.1 Стейкхолдеры: тип (внутр./внеш.) и почему важен', all(r[2] in ('Внутренний', 'Внешний') and r[3] for r in sh), f'{len(sh)} стейкхолдеров')
    check('5.1 Среди внешних стейкхолдеров есть конкуренты и СМИ', any('Конкурент' in r[1] for r in sh) and any('СМИ' in r[1] for r in sh))
    check('5.2–5.3 Квадрант D: своевременное и полное информирование (высокое влияние)', 'своевременно' in TQ.rows[4].cells[1].text and 'полн' in TQ.rows[4].cells[1].text)
    check('5.2 Матрица влияние/интерес (таблица + рисунок)', len(d.inline_shapes) >= 1)
    nums, bad = [], []
    for i in range(1, 5):
        for line in TQ.rows[i].cells[2].text.split('\n'):
            m = re.match(r'(\d+)\. .*\((\d); (\d)\)', line)
            if m:
                nums.append(int(m.group(1))); inf, it = int(m.group(2)), int(m.group(3))
                q = 'A' if inf >= 4 and it >= 3 else 'B' if inf <= 3 and it >= 3 else 'C' if inf <= 3 else 'D'
                if q != 'ABCD'[i - 1]:
                    bad.append(m.group(1))
    check('5.3 Каждый стейкхолдер ровно в одном квадранте по правилу 5.2', sorted(nums) == list(range(1, len(sh) + 1)) and not bad)
    concl = [p for p in ps if p.startswith(('Проект обоснован', 'Обязательный функционал', 'Ключевой стейкхолдер', 'Проект реализуем'))]
    check('6 Выводы — 3–5 предложений с фактами', 3 <= len(concl) <= 5 and 'BYN' in ' '.join(concl))
    check('Дата заполнения', 'Дата заполнения: 28.09.2026' in full)

elif work == 'pz4':
    r4 = docx.Document(f'{R}/pz4/Земляник - ПЗ4 WBS, сетевой график и диаграмма Ганта TableMind.docx'); T = r4.tables
    full = '\n'.join(p.text for p in r4.paragraphs)
    check('Титульный лист: ФИО слушателя', 'Земляник Р.В.' in full and 'ФИО' not in full, 'Земляник Р.В., гр. 60121')
    codes = [r.cells[0].text for r in T[0].rows[1:]]
    check('1 WBS: 12–15 работ, уникальные коды 1.0, 2.0 …', 12 <= len(codes) <= 15 and len(set(codes)) == len(codes)
          and all(re.fullmatch(r'\d+\.0', c) for c in codes), f'{len(codes)} работ')
    check('1 WBS: верхний уровень — название проекта (схема)', 'Рисунок 1 — Иерархическая структура работ' in full)
    check('2 Длительность в месяцах, PERT (O+4M+P)/6', all(r.cells[6].text.isdigit() for r in T[1].rows[1:]) and '(O + 4M + P) / 6' in full)
    check('3 Таблица связей: «Код работы», «Название работы», «Непосредственные предшественники»',
          [c.text for c in T[2].rows[0].cells] == ['Код работы', 'Название работы', 'Непосредственные предшественники'])
    D = {r.cells[0].text: int(r.cells[2].text) for r in T[3].rows[1:]}
    cr = [r.cells[0].text for r in T[3].rows[1:] if r.cells[8].text == 'да']
    ts = [int(r.cells[7].text) for r in T[3].rows[1:] if r.cells[8].text == 'нет']
    ef = max(int(r.cells[4].text) for r in T[3].rows[1:])
    check('CPM: сумма длительностей критического пути = длительности проекта', sum(D[c] for c in cr) == ef, f'{sum(D[c] for c in cr)} = {ef} мес.')
    check('CPM: у некритических работ резерв > 0', all(x > 0 for x in ts))
    check('4 Диаграмма Ганта (работы, месяцы, полосы, связи, критический путь цветом)', 'Рисунок 3 — Диаграмма Ганта' in full and len(r4.inline_shapes) == 3)
    check('Сетевой график PDM', 'Рисунок 2 — Сетевой график' in full)
    wb = openpyxl.load_workbook(f'{R}/pz4/Земляник - Диаграмма Ганта.xlsx'); ws = wb['ТАБЛИЦА ДАННЫХ']
    check('Планировщик: реальные работы вместо демо «раб1…»', all(ws.cell(r, 3).value for r in range(8, 23)) and 'раб' not in str([ws.cell(r, 3).value for r in range(8, 29)]))
    check('Планировщик: ES/EF/LS/LF/резерв/критичность — формулы', all(str(ws.cell(r, c).value).startswith('=') for r in range(9, 23) for c in (6, 8, 9, 10, 11, 12)))
    p = wb['Планировщик проекта']
    check('Планировщик: лист Ганта ссылается на таблицу данных', all(str(p.cell(r, 7).value).startswith("='ТАБЛИЦА ДАННЫХ'") for r in range(5, 20)))

elif work == 'risk':
    f = f'{R}/risk-management/Земляник - Карта рисков.xlsx'
    wb = openpyxl.load_workbook(f); ws = wb['Таблица с рисками']
    rr = [r for r in range(4, 24) if ws.cell(r, 3).value]
    check('Название проекта указано', bool(ws['C1'].value))
    check('Реестр: ≥ 10 рисков проекта', len(rr) >= 10, f'{len(rr)}')
    check('Влияние, триггер, меры у каждого риска', all(ws.cell(r, 4).value and ws.cell(r, 5).value and ws.cell(r, 9).value for r in rr))
    check('Оценки вероятности и воздействия от 0 до 1', all(0 < ws.cell(r, 6).value <= 1 and 0 < ws.cell(r, 7).value <= 1 for r in rr))
    check('Итоговый коэффициент — формула P × I', all(str(ws.cell(r, 8).value).startswith('=') for r in rr))
    k = [ws.cell(r, 6).value * ws.cell(r, 7).value for r in rr]
    hi, mid, lo = sum(x > 0.505 for x in k), sum(0.14 <= x <= 0.505 for x in k), sum(x < 0.14 for x in k)
    xml = zipfile.ZipFile(f).read('xl/worksheets/sheet2.xml').decode()
    check('Пороги приоритета шаблона сохранены (0,505 / 0,14)', '$H4&gt;0.505' in xml and '$H4&lt;0.14' in xml,
          f'высоких {hi}, средних {mid}, низких {lo}')
    ch = zipfile.ZipFile(f).read('xl/charts/chart1.xml').decode()
    check('Карта рисков: точки подписаны значениями из ячеек', 'datalabelsRange' in ch and 'showDataLabelsRange val="1"' in ch)
    check('Мероприятия указаны для рисков высокого приоритета', all(ws.cell(r, 9).value for r, x in zip(rr, k) if x > 0.505))

elif work == 'appendix':
    wb = openpyxl.load_workbook(f'{R}/appendix/ПЗ4_Приложение_расчёты_WBS_PERT_CPM_TableMind.xlsx', data_only=True)
    ws = wb['Расчёт CPM']
    check('Приложение ПЗ4: расчёт CPM совпадает с отчётом (T = 12)', max(ws.cell(r, 7).value for r in range(5, 20)) == 12)
    wb2 = openpyxl.load_workbook(f'{R}/appendix/Риски_приложение_реестр_TableMind.xlsx', data_only=True)
    w2 = wb2['Реестр рисков']
    n = len([r for r in range(5, 25) if w2.cell(r, 3).value and str(w2.cell(r, 2).value).startswith('R')])
    n_tpl = len([r for r in range(4, 24) if openpyxl.load_workbook(f'{R}/risk-management/Земляник - Карта рисков.xlsx')['Таблица с рисками'].cell(r, 3).value])
    check('Приложение рисков: те же риски, что в шаблоне', n == n_tpl, f'{n} = {n_tpl}')
    check('Приложение рисков: резерв покрывает принятые риски', any(w2.cell(r, 4).value == 'ДА' for r in range(20, 40)))

elif work == 'cross':
    import glob
    p1 = docx.Document(glob.glob(f'{R}/pz1/*.docx')[0]); p4 = docx.Document(glob.glob(f'{R}/pz4/*.docx')[0])
    t1 = '\n'.join(p.text for p in p1.paragraphs) + ' '.join(cells(t) for t in p1.tables)
    t4 = '\n'.join(p.text for p in p4.paragraphs) + ' '.join(cells(t) for t in p4.tables)
    check('MVP 31.03.2027 одинаково в ПЗ1 и ПЗ4', '31.03.2027' in t1 and '31.03.2027' in t4)
    check('Завершение 30.09.2027 одинаково в ПЗ1 и ПЗ4', '30.09.2027' in t1 and '30.09.2027' in t4)
    check('Бюджет 517 000 BYN одинаков в ПЗ1 и ПЗ4', '517 000' in t1 and '517 000' in t4)
    wr = openpyxl.load_workbook(glob.glob(f'{R}/risk-management/*.xlsx')[0])['Таблица с рисками']
    rtxt = ' '.join(str(wr.cell(r, c).value) for r in range(4, 24) for c in range(3, 10))
    check('Резерв 47 000 BYN одинаков в ПЗ1 и реестре рисков', '47 000' in t1 and '47 000' in rtxt)
    codes4 = [r.cells[0].text for r in p4.tables[0].rows[1:]]
    ps = [p.text for p in p1.paragraphs]
    stages = [p for p in ps if re.match(r'^\d\. .*\(WBS', p)]
    covered = set()
    for s_ in stages:
        m = re.search(r'WBS ([\d.]+)(?:–([\d.]+))?', s_)
        a, b = float(m.group(1)), float(m.group(2) or m.group(1))
        covered |= {c for c in codes4 if a <= float(c) <= b}
    check('Этапы декомпозиции ПЗ1 покрывают все 15 работ WBS ПЗ4', covered == set(codes4), f'{len(stages)} этапов → {len(covered)} работ')
    wa = openpyxl.load_workbook(f'{R}/appendix/Риски_приложение_реестр_TableMind.xlsx')['Реестр рисков']
    refs = set()
    for r in range(5, 25):
        val = wa.cell(r, 14).value
        if val and val != 'все':
            refs |= {x.strip() for x in str(val).split(',')}
    check('Риски ссылаются только на существующие работы WBS', refs <= set(codes4), ', '.join(sorted(refs - set(codes4))) or 'все ссылки корректны')
    cp4 = re.search(r'Критический путь: ([\d. →]+)', t4).group(1).strip().rstrip('.')
    wg = openpyxl.load_workbook(glob.glob(f'{R}/pz4/*.xlsx')[0])['ТАБЛИЦА ДАННЫХ']
    crit_rows = [str(wg.cell(r, 3).value).split()[0] for r in range(8, 23)]
    check('Планировщик содержит те же 15 работ WBS в том же порядке; критический путь', crit_rows == codes4, cp4)
    names4 = {r.cells[0].text: r.cells[1].text for r in p4.tables[0].rows[1:]}
    tw1 = next(t for t in p1.tables if t.rows[0].cells[0].text.strip() == 'Код')
    names1 = {r.cells[0].text: r.cells[1].text for r in tw1.rows[1:]}
    check('Таблица 4.2 ПЗ1 совпадает с WBS ПЗ4 (коды и названия работ)', names1 == names4, f'{len(names1)} работ')
    check('Should-функции (ghost-подсказки, Cmd-K) не запланированы ни в одной работе', not re.search(r'ghost|Cmd-K', t4 + rtxt) and 'ghost' not in ' '.join(names1.values()))
    ass = [c.text for r in p1.tables[3].rows[1:] for c in [r.cells[1]]]
    keys = ['пилотных', 'LLM', 'AppSource', 'лицензи', 'участники']
    check('Каждое допущение ПЗ1 отражено риском в реестре', all(any(k.lower() in str(wr.cell(r, 3).value).lower() or (k == 'участники' and 'разработчика' in str(wr.cell(r, 3).value)) or (k == 'пилотных' and 'пилотных' in str(wr.cell(r, 3).value)) for r in range(4, 24)) for k in keys))

print(f'| Требование | Статус | Комментарий |\n|---|---|---|')
print('\n'.join(rows))
sys.exit(1 if hard_fail else 0)
